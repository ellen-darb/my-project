"""Чтение исходных файлов ментора в аккуратные таблицы.

Вход: каталог с распакованным архивом «Данные Сириус» (путь аргументом).
Выход: data/entries_15min.csv.gz (дата, время, вестибюль, входы)
       data/entries_hourly.csv.gz (дата-час, вестибюль, входы)
"""
import sys, re, glob, os, datetime as dt
import openpyxl, pandas as pd

SRC = sys.argv[1]
OUT = os.path.join(os.path.dirname(__file__), "..", "data")
os.makedirs(OUT, exist_ok=True)

rows = []
for f in sorted(glob.glob(os.path.join(SRC, "*", "*по 15-мин.xlsx"))):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for name in wb.sheetnames:
        m = re.search(r"(\d{2})(\d{2})(\d{4})", name)
        if not m:
            continue
        day = dt.date(int(m[3]), int(m[2]), int(m[1]))
        it = wb[name].iter_rows(values_only=True)
        next(it)
        times = next(it)[2:]
        next(it)                                  # строка итогов по слотам
        for r in it:
            if not r or not isinstance(r[0], str) or r[0].startswith(("Итого", "Всего")):
                continue
            for t, v in zip(times, r[2:]):
                if isinstance(t, dt.time) and v is not None:
                    rows.append((day, t.hour * 60 + t.minute, r[0].strip(), float(v)))
q = pd.DataFrame(rows, columns=["date", "minute", "vestibule", "entries"])
q["date"] = pd.to_datetime(q["date"])
q.to_csv(os.path.join(OUT, "entries_15min.csv.gz"), index=False)
print("15 мин:", len(q), "строк,", q.date.nunique(), "дней,", q.vestibule.nunique(), "вестибюлей")

h = pd.read_excel(glob.glob(os.path.join(SRC, "*", "Пассажиропоток 2026*.xlsx"))[0],
                  header=None, skiprows=7, names=["dt", "vestibule", "entries"])
h = h[h.vestibule.notna()].copy()
h["ts"] = pd.to_datetime(h.dt, format="%d.%m.%Y %H")
h = h[["ts", "vestibule", "entries"]]
h.to_csv(os.path.join(OUT, "entries_hourly.csv.gz"), index=False)
print("часовые:", len(h), "строк,", h.ts.dt.date.nunique(), "дней", h.ts.min(), h.ts.max())
